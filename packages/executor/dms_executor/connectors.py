"""CONNECT-01 (dms#366): read-only source connectors behind one interface.

Every source answers the same three calls - ``introspect_schema()``,
``sample(table, n)`` and ``run_select(sql)`` - for PostgreSQL, MySQL, SQL Server,
SQLite, DuckDB, CSV, Excel and Snowflake. Swap scenario (hard rule 6): the source kind
is the swap. A customer moving MySQL -> Postgres changes ``ConnectionRef.kind``, not
callers. This is not a sixth port. Secrets go through the existing ``SecretsPort``.

Default OFF. ``open_connector`` refuses unless ``DMS_CONNECTORS_ENABLED=1``, and no
API route or Studio surface imports this module.

DR-0005 is extract-only, and ``run_select`` sends a caller's SELECT into the source
engine. So this layer is a library and is not wired into the ask path. Wiring it there
needs the P-DMS-28 unlock, and this module does not fire it.

Credentials: a ``ConnectionRef`` carries a ``service_id``, never a secret. The secret
is resolved from OpenVault at call time, used for one connection, and dropped. It is
never stored on the connector, logged, or put in an exception.

Guard: ``guard_select`` parses with sqlglot (already the executor's SQL parser, see
``sql_currency.py``). It allows one SELECT/WITH statement and refuses comments,
DDL/DML, ``SELECT INTO``, locking reads and file/network/side-effect functions, each
with a named ``SelectRejected.code``. Each engine also opens read-only where it can,
so a guard miss still meets a read-only session.
"""

from __future__ import annotations

import csv
import logging
import os
import re
import sqlite3
import tempfile
import threading
import time
from collections.abc import Iterator, Mapping
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

import httpx
import sqlglot
from dms_core.ports import SecretsPort
from sqlglot import exp
from sqlglot.errors import ParseError, TokenError

from dms_executor import db_connector as dbc

logger = logging.getLogger(__name__)

FLAG = "DMS_CONNECTORS_ENABLED"

ConnectorKind = Literal[
    "postgresql", "mysql", "sqlserver", "sqlite", "duckdb", "csv", "excel", "snowflake"
]
SERVER_KINDS = frozenset({"postgresql", "mysql", "sqlserver", "snowflake"})
FILE_KINDS = frozenset({"sqlite", "duckdb", "csv", "excel"})

DEFAULT_MAX_ROWS = 1_000
#: Hard ceiling on ``max_rows``: results sit in memory, so the cap is not optional.
MAX_ROWS_CEILING = 10_000
DEFAULT_TIMEOUT_S = 30.0

_DIALECT: dict[str, str] = {
    "postgresql": "postgres",
    "mysql": "mysql",
    "sqlserver": "tsql",
    "sqlite": "sqlite",
    "duckdb": "duckdb",
    "csv": "duckdb",
    "excel": "duckdb",
    "snowflake": "snowflake",
}

_CREDENTIAL_KEY = re.compile(r"pass|pwd|secret|token|key|credential", re.IGNORECASE)


class ConnectorError(RuntimeError):
    """A connector call failed. Messages never carry a resolved secret."""


class ConnectorsDisabled(ConnectorError):
    """``DMS_CONNECTORS_ENABLED`` is not on."""


class SecretUnavailable(ConnectorError):
    """OpenVault did not return a secret for the ``service_id``."""


class SelectTimeout(ConnectorError):
    """The SELECT ran past the connector's timeout and was cancelled."""


class SelectRejected(ConnectorError):
    """The guard refused the SQL. ``code`` names the reason."""

    def __init__(self, code: str, detail: str = "") -> None:
        self.code = code
        super().__init__(f"{code}: {detail}" if detail else code)


def connectors_enabled() -> bool:
    return os.environ.get(FLAG, "0").strip().lower() in {"1", "true", "yes", "on"}


# --- guard -------------------------------------------------------------------

_FORBIDDEN_NODES: tuple[type[exp.Expression], ...] = tuple(
    getattr(exp, name)
    for name in (
        "Insert", "Update", "Delete", "Merge", "Create", "Drop", "Alter", "TruncateTable",
        "Command", "Copy", "Pragma", "Use", "Set", "Transaction", "Commit", "Rollback",
        "LoadData", "Grant", "Refresh", "Cache", "Uncache", "Kill", "Analyze", "Describe",
    )
    if hasattr(exp, name)
)

#: File, network, extension and side-effect functions. A read-only session does not
#: stop most of these (``read_csv`` reads the host disk, ``nextval`` mutates).
_FORBIDDEN_FUNCTIONS = frozenset({
    # DuckDB file / remote readers
    "read_csv", "read_csv_auto", "read_parquet", "parquet_scan", "read_json",
    "read_json_auto", "read_json_objects", "read_ndjson", "read_text", "read_blob",
    "glob", "sniff_csv", "query", "query_table", "postgres_scan", "sqlite_scan",
    # PostgreSQL
    "pg_read_file", "pg_read_binary_file", "pg_ls_dir", "pg_stat_file", "lo_import",
    "lo_export", "dblink", "dblink_exec", "pg_terminate_backend", "pg_cancel_backend",
    "set_config", "pg_reload_conf", "nextval", "setval",
    # MySQL
    "load_file", "sys_exec", "sys_eval", "get_lock", "release_lock",
    # SQL Server
    "openrowset", "opendatasource", "openquery", "openxml",
    # SQLite
    "load_extension", "readfile", "writefile", "edit",
})


def _function_name(node: exp.Func) -> str:
    raw = node.name if isinstance(node, exp.Anonymous) else node.sql_name()
    return str(raw).lower()


def guard_select(sql: str, kind: str) -> str:
    """Return ``sql`` (trailing ``;`` stripped) if it is one read-only SELECT/WITH.

    ponytail: the guard trusts sqlglot's read of the dialect. The parser-vs-engine gap
    is closed by the read-only session each engine opens, not by more patterns here.
    """
    dialect = _DIALECT[kind]
    text = (sql or "").strip()
    if not text:
        raise SelectRejected("empty_sql")
    try:
        tokens = sqlglot.Dialect.get_or_raise(dialect).tokenize(text)
    except TokenError as exc:
        raise SelectRejected("parse_error", type(exc).__name__) from None
    if any(t.comments for t in tokens):
        raise SelectRejected("comment_smuggling", "comments are not allowed")
    try:
        statements = [s for s in sqlglot.parse(text, read=dialect) if s is not None]
    except ParseError as exc:
        raise SelectRejected("parse_error", type(exc).__name__) from None
    if not statements:
        raise SelectRejected("empty_sql")
    if len(statements) > 1:
        raise SelectRejected("multi_statement", f"{len(statements)} statements")
    root = statements[0]
    if not isinstance(root, (exp.Select, exp.SetOperation, exp.Subquery)):
        raise SelectRejected("not_select", type(root).__name__)
    for node in root.walk():
        if isinstance(node, _FORBIDDEN_NODES):
            raise SelectRejected("ddl_dml", type(node).__name__)
        if isinstance(node, exp.Into):
            raise SelectRejected("select_into")
        if isinstance(node, exp.Lock):
            raise SelectRejected("locking_read")
        if isinstance(node, exp.Func):
            name = _function_name(node)
            if name in _FORBIDDEN_FUNCTIONS or name.startswith("system$"):
                raise SelectRejected("forbidden_function", name)
    return text.rstrip(";").rstrip()


# --- secrets -----------------------------------------------------------------


class OpenVaultSecrets:
    """``SecretsPort`` over OpenVault's existing ``POST /keys/services`` call.

    Same request ``ManifestMinter.fetch_intermediate`` makes (``service_id`` in, token
    out). Unauthenticated: no admin token is sent.
    """

    def __init__(
        self, *, openvault_url: str | None = None, http: httpx.Client | None = None
    ) -> None:
        self.openvault_url = (
            openvault_url or os.environ.get("OPENVAULT_URL", "http://127.0.0.1:5000")
        ).rstrip("/")
        self._http = http

    def get(self, name: str) -> str | None:
        http = self._http or httpx.Client(timeout=10.0)
        try:
            resp = http.post(
                f"{self.openvault_url}/keys/services",
                json={"service_id": name},
                headers={"X-OpenVault-Reveal": "intentional"},
            )
        except httpx.HTTPError as exc:
            raise SecretUnavailable(
                f"openvault unreachable for {name}: {type(exc).__name__}"
            ) from None
        finally:
            if self._http is None:
                http.close()
        if resp.status_code == 404:
            return None
        if resp.status_code >= 400:
            raise SecretUnavailable(f"openvault HTTP {resp.status_code} for {name}")
        token = resp.json().get("token")
        return str(token) if token else None


# --- connection reference ----------------------------------------------------


@dataclass(frozen=True)
class ConnectionRef:
    """Where a source is and which OV ``service_id`` holds its secret. No secret.

    ``options`` takes non-secret knobs only (Snowflake ``account``/``warehouse``/
    ``role``, SQL Server ``trust_server_certificate``). A key that looks like a
    credential is refused.
    """

    kind: ConnectorKind
    service_id: str = ""
    host: str = ""
    port: int | None = None
    database: str = ""
    user: str = ""
    path: str = ""
    options: Mapping[str, str] = field(default_factory=dict)

    def __post_init__(self) -> None:
        if self.kind not in _DIALECT:
            raise ValueError(f"unsupported connector kind: {self.kind!r}")
        bad = sorted(k for k in self.options if _CREDENTIAL_KEY.search(k))
        if bad:
            raise ValueError(f"credentials come only from OpenVault; refused options {bad}")
        if self.kind in SERVER_KINDS and not self.service_id:
            raise ValueError(f"{self.kind} needs an OpenVault service_id")
        if self.kind in FILE_KINDS and not self.path:
            raise ValueError(f"{self.kind} needs a path")

    def describe(self) -> str:
        """Credential-free label, safe to log and display."""
        if self.kind in FILE_KINDS:
            return f"{self.kind}://{Path(self.path).name}"
        if self.kind == "snowflake":
            return f"snowflake://{self.options.get('account', '')}/{self.database}"
        return f"{self.kind}://{self.host}:{self.port or ''}/{self.database}"


# --- result shapes -----------------------------------------------------------


@dataclass(frozen=True)
class ColumnInfo:
    name: str
    type: str


@dataclass(frozen=True)
class TableInfo:
    schema: str
    name: str
    columns: tuple[ColumnInfo, ...]

    @property
    def qualified(self) -> str:
        return f"{self.schema}.{self.name}"


@dataclass(frozen=True)
class SelectResult:
    columns: list[str]
    rows: list[tuple[Any, ...]]
    #: True when the source had more rows than ``row_cap``.
    truncated: bool
    row_cap: int


# --- the connector -----------------------------------------------------------

_SYSTEM_SCHEMAS: dict[str, set[str]] = {
    **dbc._SYSTEM_SCHEMAS,
    "duckdb": {"information_schema", "pg_catalog"},
    "csv": {"information_schema", "pg_catalog"},
    "excel": {"information_schema", "pg_catalog"},
    "snowflake": {"INFORMATION_SCHEMA"},
}

_COLUMNS_SQL = (
    "SELECT TABLE_SCHEMA, TABLE_NAME, COLUMN_NAME, DATA_TYPE "
    "FROM INFORMATION_SCHEMA.COLUMNS "
)
_COLUMNS_ORDER = "ORDER BY TABLE_SCHEMA, TABLE_NAME, ORDINAL_POSITION"


def _safe_stem(name: str) -> str:
    stem = "".join(c if c.isalnum() else "_" for c in name).strip("_").lower()
    return stem[:60] or "sheet"


def _scrub(text: str, secret: str | None) -> str:
    return text.replace(secret, "[REDACTED]") if secret else text


def _is_timeout(exc: BaseException) -> bool:
    name = type(exc).__name__.lower()
    text = str(exc).lower()
    return (
        "interrupt" in name
        or "querycanceled" in name
        or any(
            s in text
            for s in (
                "statement timeout", "canceling statement", "maximum statement execution",
                "query timeout", "hyt00", "interrupted", "timeout expired",
            )
        )
    )


class ReadOnlyConnector:
    """``introspect_schema()``, ``sample(table, n)``, ``run_select(sql)`` for one source.

    Opens one connection per call and resolves the secret inside it.
    ponytail: CSV/Excel reload into an in-memory DuckDB on every call; fine at SME file
    sizes, cache the loaded handle if a measured file makes that slow.
    """

    def __init__(
        self,
        ref: ConnectionRef,
        *,
        secrets: SecretsPort | None = None,
        max_rows: int = DEFAULT_MAX_ROWS,
        timeout_s: float = DEFAULT_TIMEOUT_S,
    ) -> None:
        if max_rows < 1:
            raise ValueError("max_rows must be >= 1")
        if timeout_s <= 0:
            raise ValueError("timeout_s must be > 0")
        self.ref = ref
        self.max_rows = min(max_rows, MAX_ROWS_CEILING)
        self.timeout_s = float(timeout_s)
        self._secrets: SecretsPort = secrets or OpenVaultSecrets()

    def __repr__(self) -> str:
        return f"ReadOnlyConnector({self.ref.describe()})"

    # -- the interface --

    def introspect_schema(self) -> list[TableInfo]:
        with self._session() as (con, secret):
            return self._introspect(con, secret)

    def sample(self, table: str, n: int = 10) -> SelectResult:
        cap = max(1, min(int(n), self.max_rows))
        with self._session() as (con, secret):
            target = self._resolve(self._introspect(con, secret), table)
            ident = ".".join(self._quote(p) for p in (target.schema, target.name))
            if self.ref.kind == "sqlserver":
                sql = f"SELECT TOP ({cap}) * FROM {ident}"
            else:
                sql = f"SELECT * FROM {ident} LIMIT {cap}"
            return self._select(con, secret, guard_select(sql, self.ref.kind), cap)

    def run_select(self, sql: str) -> SelectResult:
        checked = guard_select(sql, self.ref.kind)
        with self._session() as (con, secret):
            return self._select(con, secret, checked, self.max_rows)

    # -- internals --

    def _quote(self, part: str) -> str:
        if self.ref.kind == "sqlserver":
            return "[" + part.replace("]", "]]") + "]"
        if self.ref.kind == "mysql":
            return "`" + part.replace("`", "``") + "`"
        return '"' + part.replace('"', '""') + '"'

    def _resolve(self, tables: list[TableInfo], wanted: str) -> TableInfo:
        """Identifiers come from the source's own catalog, never the caller's string."""
        schema, _, name = wanted.rpartition(".")
        hits = [
            t for t in tables
            if t.name.lower() == name.lower() and (not schema or t.schema.lower() == schema.lower())
        ]
        if len(hits) != 1:
            why = "ambiguous; qualify the schema" if hits else "not found"
            raise dbc.UnknownSourceTable(f"{wanted!r} {why} on {self.ref.describe()}")
        return hits[0]

    def _introspect(self, con: Any, secret: str | None) -> list[TableInfo]:
        kind = self.ref.kind
        if kind == "sqlite":
            return self._introspect_sqlite(con)
        sql = _COLUMNS_SQL
        params: tuple[str, ...] = ()
        if kind == "mysql":
            sql, params = sql + "WHERE TABLE_SCHEMA = %s ", (self.ref.database,)
        rows = self._fetch_all(con, secret, sql + _COLUMNS_ORDER, params)
        skip = _SYSTEM_SCHEMAS[kind]
        tables: dict[tuple[str, str], list[ColumnInfo]] = {}
        for schema, table, column, dtype in rows:
            if str(schema) in skip:
                continue
            tables.setdefault((str(schema), str(table)), []).append(
                ColumnInfo(str(column), str(dtype))
            )
        return [TableInfo(s, t, tuple(cols)) for (s, t), cols in tables.items()]

    def _introspect_sqlite(self, con: sqlite3.Connection) -> list[TableInfo]:
        names = [
            r[0]
            for r in con.execute(
                "SELECT name FROM sqlite_master WHERE type IN ('table', 'view') "
                "AND name NOT LIKE 'sqlite_%' ORDER BY name"
            ).fetchall()
        ]
        return [
            TableInfo(
                "main",
                name,
                tuple(
                    ColumnInfo(str(c[0]), str(c[1]))
                    for c in con.execute(
                        "SELECT name, type FROM pragma_table_info(?) ORDER BY cid", (name,)
                    ).fetchall()
                ),
            )
            for name in names
        ]

    def _fetch_all(
        self, con: Any, secret: str | None, sql: str, params: tuple[Any, ...]
    ) -> list[tuple[Any, ...]]:
        with self._cursor(con, secret) as cur:
            if params:
                cur.execute(sql, params)
            else:
                cur.execute(sql)
            return [tuple(r) for r in cur.fetchall()]

    def _select(self, con: Any, secret: str | None, sql: str, cap: int) -> SelectResult:
        started = time.monotonic()
        with self._cursor(con, secret) as cur:
            cur.execute(sql)
            columns = [str(d[0]) for d in (cur.description or [])]
            fetched = [tuple(r) for r in cur.fetchmany(cap + 1)]
        # Some engines cancel softly: MySQL's MAX_EXECUTION_TIME stops BENCHMARK() and
        # returns 0 with no error. A result that ran into the deadline is not an answer.
        if time.monotonic() - started >= self.timeout_s:
            raise SelectTimeout(f"{self.ref.describe()}: ran past {self.timeout_s:g}s")
        return SelectResult(columns, fetched[:cap], len(fetched) > cap, cap)

    @contextmanager
    def _cursor(self, con: Any, secret: str | None) -> Iterator[Any]:
        timer: threading.Timer | None = None
        deadline = time.monotonic() + self.timeout_s
        cur = con.cursor()
        if self.ref.kind in {"duckdb", "csv", "excel"}:
            # A DuckDB cursor is its own connection: interrupt the cursor, not ``con``.
            timer = threading.Timer(self.timeout_s, cur.interrupt)
            timer.start()
        elif self.ref.kind == "sqlite":
            con.set_progress_handler(lambda: int(time.monotonic() > deadline), 1_000)
        try:
            yield cur
        except (ConnectorError, dbc.UnknownSourceTable):
            raise
        except Exception as exc:
            if _is_timeout(exc):
                raise SelectTimeout(
                    f"{self.ref.describe()}: cancelled after {self.timeout_s:g}s"
                ) from None
            detail = _scrub(f"{type(exc).__name__}: {exc}", secret)
            raise ConnectorError(f"{self.ref.describe()}: {detail}") from None
        finally:
            if timer is not None:
                timer.cancel()
            try:
                cur.close()
            except Exception:  # noqa: BLE001 - close must not mask the real error
                pass

    def _resolve_secret(self) -> str:
        secret = self._secrets.get(self.ref.service_id)
        if not secret:
            raise SecretUnavailable(f"no secret in OpenVault for {self.ref.service_id}")
        return secret

    @contextmanager
    def _session(self) -> Iterator[tuple[Any, str | None]]:
        kind = self.ref.kind
        if kind in FILE_KINDS:
            with self._open_file() as con:
                yield con, None
            return
        secret = self._resolve_secret()
        try:
            opened = self._open_server(secret)
            con = opened.__enter__()
        except ConnectorError:
            raise
        except Exception as exc:
            cause = exc.__cause__ or exc
            detail = _scrub(f"{type(cause).__name__}: {cause}", secret)
            raise ConnectorError(f"could not connect to {self.ref.describe()}: {detail}") from None
        logger.info("connector session open %s", self.ref.describe())
        try:
            try:
                self._prepare_server(con)
            except Exception as exc:
                detail = _scrub(f"{type(exc).__name__}: {exc}", secret)
                raise ConnectorError(f"{self.ref.describe()}: {detail}") from None
            yield con, secret
        finally:
            opened.__exit__(None, None, None)
            del secret

    def _open_server(self, secret: str) -> Any:
        ref = self.ref
        if ref.kind == "snowflake":
            return self._open_snowflake(secret)
        cfg = dbc.SourceConfig(
            kind=ref.kind,  # type: ignore[arg-type]
            host=ref.host,
            database=ref.database,
            user=ref.user,
            password=secret,
            port=ref.port,
            encrypt=ref.options.get("encrypt", "true").lower() == "true",
            trust_server_certificate=(
                ref.options.get("trust_server_certificate", "false").lower() == "true"
            ),
            connect_timeout=max(1, int(self.timeout_s)),
        )
        return dbc.connect(cfg)

    @contextmanager
    def _open_snowflake(self, secret: str) -> Iterator[Any]:
        try:
            import snowflake.connector as sf  # type: ignore[import-not-found]
        except ImportError:
            raise ConnectorError(
                "snowflake-connector-python is not installed; Snowflake is unavailable"
            ) from None
        opts = self.ref.options
        con = sf.connect(
            account=opts.get("account", ""),
            user=self.ref.user,
            password=secret,
            database=self.ref.database,
            warehouse=opts.get("warehouse") or None,
            role=opts.get("role") or None,
            login_timeout=max(1, int(self.timeout_s)),
            session_parameters={
                "STATEMENT_TIMEOUT_IN_SECONDS": max(1, int(self.timeout_s)),
                "QUERY_TAG": "dms-connect-01",
            },
        )
        try:
            yield con
        finally:
            con.close()

    def _prepare_server(self, con: Any) -> None:
        """Make the session read-only and time-boxed on the engine's side too."""
        ms = max(1, int(self.timeout_s * 1000))
        kind = self.ref.kind
        if kind == "postgresql":
            con.autocommit = True
            con.execute("SET default_transaction_read_only = on")
            con.execute(f"SET statement_timeout = {ms}")
        elif kind == "mysql":
            con.autocommit(True)
            cur = con.cursor()
            try:
                cur.execute("SET SESSION TRANSACTION READ ONLY")
                cur.execute(f"SET SESSION MAX_EXECUTION_TIME = {ms}")
            finally:
                cur.close()
        elif kind == "sqlserver":
            # pyodbc query timeout; the connection is already opened readonly=True.
            con.timeout = max(1, int(self.timeout_s))

    @contextmanager
    def _open_file(self) -> Iterator[Any]:
        import duckdb

        path = Path(self.ref.path)
        if not path.is_file():
            raise ConnectorError(f"{self.ref.describe()}: file not found")
        kind = self.ref.kind
        if kind == "sqlite":
            con = sqlite3.connect(f"{path.resolve().as_uri()}?mode=ro", uri=True)
            con.set_authorizer(_sqlite_read_only)
            try:
                yield con
            finally:
                con.close()
            return
        if kind == "duckdb":
            dcon = duckdb.connect(
                str(path), read_only=True, config={"enable_external_access": False}
            )
        else:
            dcon = duckdb.connect()
            if kind == "csv":
                dcon.read_csv(str(path)).create(_safe_stem(path.stem))
            else:
                _load_excel(dcon, path)
            dcon.execute("SET enable_external_access = false")
            dcon.execute("SET lock_configuration = true")
        try:
            yield dcon
        finally:
            dcon.close()


_SQLITE_ALLOWED = frozenset({
    sqlite3.SQLITE_SELECT, sqlite3.SQLITE_READ, sqlite3.SQLITE_FUNCTION,
    getattr(sqlite3, "SQLITE_RECURSIVE", 33),
})


def _sqlite_read_only(action: int, arg1: Any, *_: Any) -> int:
    if action in _SQLITE_ALLOWED:
        return sqlite3.SQLITE_OK
    if action == sqlite3.SQLITE_PRAGMA and arg1 == "table_info":
        return sqlite3.SQLITE_OK
    return sqlite3.SQLITE_DENY


def _load_excel(con: Any, path: Path) -> None:
    """Each sheet becomes one table. The workbook is opened read-only and never saved."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        with tempfile.TemporaryDirectory() as tmp:
            for sheet in wb.worksheets:
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue
                out = Path(tmp) / f"{_safe_stem(sheet.title)}.csv"
                with out.open("w", newline="", encoding="utf-8") as fh:
                    csv.writer(fh).writerows(
                        ["" if v is None else v for v in row] for row in rows
                    )
                con.read_csv(str(out), header=True).create(_safe_stem(sheet.title))
    finally:
        wb.close()


def open_connector(
    ref: ConnectionRef,
    *,
    secrets: SecretsPort | None = None,
    max_rows: int = DEFAULT_MAX_ROWS,
    timeout_s: float = DEFAULT_TIMEOUT_S,
) -> ReadOnlyConnector:
    """The only public way in. Refuses while ``DMS_CONNECTORS_ENABLED`` is off."""
    if not connectors_enabled():
        raise ConnectorsDisabled(f"{FLAG} is off; connectors are disabled by default")
    return ReadOnlyConnector(ref, secrets=secrets, max_rows=max_rows, timeout_s=timeout_s)

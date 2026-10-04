"""Plan E gate matrix, theme "mask-envelope": answer-path text that skips the PII masker (G5).

The envelope masker (``dms_core.pii``) is applied to ``text``, ``rows``, ``values``,
``sql_used`` and, through ``mask_unknown_keys``, to ``assumptions`` and the other
loose keys. Three places sit outside that walk today. Each test below shows one of
them on the user-visible output, after a control that proves the masker ran on the
SAME request (so the failing assertion is this gap and not a broken fixture):

* ``audit_receipt`` (``exclude.reasons[].detail``, ``unsure.why``) and source cards
  attached after the build (``contributing_sources`` from the CCA cascade). The
  receipt is in ``_HANDLED_KEYS`` so ``mask_unknown_keys`` skips it, and
  ``mask_envelope`` only rewrites ``include.rows``. The raw literal reaches the
  ask envelope the UI renders, the .xlsx workbook and the BI stub.
* The Cortex ledger append for verified-query and generative answers: the payload
  carries the raw SQL text while the envelope's ``sql_used`` is masked.
* Engine error text raised as ``AskServiceError``: returned verbatim as
  ``detail.message`` on the HTTP error with no masker.

Every literal is synthetic (a made-up IC shape and an example.com address). No key,
network, server or model call: an in-process Executor over a recording fake Cortex.
"""

from __future__ import annotations

import io
import json
from typing import Any

import duckdb
from _harness import (
    FINANCE,
    HOSTILE_CELLS,
    assert_envelope,
    control,
    gap,
    require_envelope,
)
from cortex_client.models import AskResponse
from openpyxl import load_workbook

#: Malaysian IC (NRIC) shape. Synthetic, not a real person. The masker tokenises it
#: as ``DMSMASK_nric_NN`` wherever it does reach it.
NRIC_RECEIPT = "900101-14-5678"
NRIC_LEDGER = "900101-14-5555"
#: Reserved example domain (see HOSTILE_CELLS["pii"]); the masker tokenises it as email.
EMAIL = HOSTILE_CELLS["pii"]


# --------------------------------------------------------------------------
# shared controls
# --------------------------------------------------------------------------


def _control_pack(h: Any) -> None:
    """Control on the already-gated path: a certified pack answer, end to end."""
    from dms_executor.demo_pack import SPEND_BY_COUNTRY_Q

    status, env = h.ask(SPEND_BY_COUNTRY_Q, space_id=FINANCE, session_id="gm_ctl")
    control(status == 200 and isinstance(env, dict), f"pack control: HTTP {status} {env!r}")
    assert_envelope(env)
    control(
        env["badge"] == "L1_GOVERNED_METRIC" and env["abstained"] is False and env["rows"],
        f"pack control did not certify: {env.get('text')!r}",
    )


def _allow_studio_gate(monkeypatch: Any) -> None:
    """Stub the F5 gate on the Studio route only (the fake Cortex has no base_url, so
    a mutation would fail closed with 403 before it reached the code under test)."""
    import dms_api.routes.studio as studio
    from cortex_client.gate import ComplianceDecision

    monkeypatch.setattr(
        studio,
        "compliance_gate",
        lambda *, action, **_: ComplianceDecision(
            allowed=True, reason="test_allow", action=action
        ),
    )


def _register_verified_query(h: Any, question: str, sql: str) -> None:
    r = h.client.post(
        "/v1/studio/verified-queries",
        json={"space_id": FINANCE, "question": question, "sql": sql},
    )
    control(r.status_code == 200, f"verified-query registration: HTTP {r.status_code} {r.text}")


def _sheet_cells(data: bytes) -> dict[str, list[str]]:
    """Every non-empty cell of every sheet, as text: what a person opening the file sees."""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: dict[str, list[str]] = {}
    for ws in wb.worksheets:
        out[ws.title] = [
            str(c) for row in ws.iter_rows(values_only=True) for c in row if c is not None
        ]
    return out


def _cover_fields(sheets: dict[str, list[str]]) -> dict[str, str]:
    """Cover is a two-column grid (field, value) flattened row by row, header first."""
    flat = sheets["Cover"]
    return {flat[i]: flat[i + 1] for i in range(2, len(flat) - 1, 2)}


def _receipt_envelope(literal: str) -> dict[str, Any]:
    """A caller envelope whose SQL and audit receipt both quote one filter literal."""
    return {
        "answer_id": "a1",
        "badge": "L2_VALIDATED",
        "text": "x",
        "sql_used": f"SELECT n FROM t WHERE ic_no = '{literal}'",
        "audit_receipt": {
            "include": {"status": "na", "rows": []},
            "exclude": {
                "status": "filters",
                "reasons": [{"kind": "where", "detail": f"ic_no = '{literal}'"}],
                "why": "w",
            },
            "unsure": {"status": "none", "why": "w"},
        },
    }


# --------------------------------------------------------------------------
# gap 1: audit_receipt and late-added source cards sit outside the masker's walk
# --------------------------------------------------------------------------


@gap(
    "chat-ask:generative-named-abstain, chat-ask:cortex-contract-ask, "
    "chat-ask:cortex-doc-retrieval",
    "G5",
    "new",
    "audit-receipt-and-late-added-cards-skip-mask",
)
def test_ask_abstain_receipt_why_is_masked_like_assumptions(harness_factory: Any) -> None:
    """The ABSTAIN reason is masked in assumptions[0] and printed raw in audit_receipt.unsure.why."""
    resp = AskResponse(
        answer="I could not match that filter to anything in this Space.",
        badge="abstain",
        abstained=True,
        audit_id="aud_gm_abstain_receipt",
        route="abstain",
        assumptions=[f"interpreted the filter as ic_no = '{NRIC_RECEIPT}'"],
    )
    h = harness_factory(ask_response=resp)
    _control_pack(h)

    status, env = h.ask(
        f"List payments for ic_no {NRIC_RECEIPT}", space_id=FINANCE, session_id="gm_receipt"
    )
    require_envelope(status, env)
    control(env["abstained"] is True and env["badge"] == "ABSTAIN", env.get("text"))
    # The masker ran on this very envelope: the same reason, in assumptions, is tokenised.
    control(
        "DMSMASK_nric" in env["assumptions"][0] and NRIC_RECEIPT not in env["assumptions"][0],
        f"assumptions were not masked, so the fixture is wrong: {env['assumptions']!r}",
    )
    control(
        env["audit_receipt"]["unsure"]["status"] == "abstain",
        f"no abstain receipt to inspect: {env['audit_receipt']!r}",
    )

    # The UI renders the receipt (AnswerMessage.tsx) and the share payload copies it whole.
    assert NRIC_RECEIPT not in json.dumps(env["audit_receipt"]), (
        "audit_receipt.unsure.why carries the raw literal the envelope masks elsewhere: "
        f"{env['audit_receipt']['unsure']['why']!r}"
    )


@gap(
    "chat-ask:cascade-attach",
    "G5",
    "new",
    "audit-receipt-and-late-added-cards-skip-mask",
)
def test_cascade_encoding_source_cards_are_masked_like_the_trace(harness_factory: Any) -> None:
    """CCA encoding cards are attached after the build: masked in the trace, raw in contributing_sources."""
    resp = AskResponse(
        answer="Lease revenue is 350.0 MYR.",
        badge="certified",
        route="sql",
        audit_id="aud_gm_cascade",
        sql_used="SELECT SUM(quantity_kg) AS total FROM transactions",
        rows=[{"total": 350.0}],
        values=[{"id": "v0", "value": 350.0, "label": "total"}],
        drillthrough_token="tok_gm_cascade",
        contributing_sources=[
            {
                "ref_id": "transactions",
                "container": "transactions",
                "kind": "sql",
                "row_count": 1,
                "snippet": "sum",
            }
        ],
    )
    h = harness_factory(ask_response=resp, cascade=True)
    _control_pack(h)

    # Land the encodings the cascade binds (asset class, tenure) on a granted table, with one
    # asset_class value that is an email address: a value the filter leaves out.
    con = duckdb.connect(str(h.warehouse))
    con.execute("ALTER TABLE transactions ADD COLUMN asset_class VARCHAR")
    con.execute("ALTER TABLE transactions ADD COLUMN transaction_type VARCHAR")
    con.execute("UPDATE transactions SET asset_class = 'COM', transaction_type = 'LEASE'")
    con.execute("UPDATE transactions SET asset_class = ? WHERE txn_id = 'T001'", [EMAIL])
    con.close()

    status, env = h.ask(
        "lease revenue across SEA for commercial property",
        space_id=FINANCE,
        session_id="gm_cascade",
    )
    require_envelope(status, env)
    trace = {c["type"]: c for c in env.get("constraint_trace") or []}
    control(
        "asset_class" in trace and trace["asset_class"]["status"] == "CERTIFIED",
        f"the cascade did not engage and certify: {env.get('constraint_trace')!r}",
    )
    # The masker ran on this envelope: the same landed value is a token in the trace and notes.
    control(
        any("DMSMASK_email" in line for line in trace["asset_class"]["evidence"]),
        f"constraint_trace was not masked, so the fixture is wrong: {trace['asset_class']!r}",
    )
    control(
        EMAIL not in json.dumps(env["assumptions"]),
        f"assumptions were not masked, so the fixture is wrong: {env['assumptions']!r}",
    )
    cards = [
        s
        for s in env["contributing_sources"]
        if str(s.get("ref_id", "")).startswith("cca_asset_class")
    ]
    control(cards, f"no CCA encoding card was attached: {env['contributing_sources']!r}")

    assert EMAIL not in json.dumps(env["contributing_sources"]), (
        "the attached encoding card carries the raw landed value the trace masks: "
        f"{cards[0].get('snippet')!r}"
    )


@gap(
    "export:xlsx, ui:share-answer",
    "G5",
    "new",
    "audit-receipt-and-late-added-cards-skip-mask",
)
def test_xlsx_export_masks_the_audit_receipt_like_sql_used(harness: Any) -> None:
    """In the downloaded workbook the Cover sql_used row is tokenised and the audit_receipt row is raw."""
    _control_pack(harness)

    r = harness.client.post(
        "/v1/chat/export.xlsx", json={"envelope": _receipt_envelope(NRIC_RECEIPT)}
    )
    control(
        r.status_code == 200 and r.content[:2] == b"PK",
        f"export.xlsx did not return a workbook: HTTP {r.status_code} {r.text[:200]}",
    )
    sheets = _sheet_cells(r.content)
    cover = _cover_fields(sheets)
    # The masker ran on this workbook: the SQL row is a token and the raw literal is not in it.
    control(
        "DMSMASK_nric" in cover.get("sql_used", "") and NRIC_RECEIPT not in cover["sql_used"],
        f"Cover sql_used was not masked, so the fixture is wrong: {cover.get('sql_used')!r}",
    )
    control("audit_receipt" in cover, f"Cover has no audit_receipt row: {sorted(cover)!r}")

    leaked = {name: [c for c in cells if NRIC_RECEIPT in c] for name, cells in sheets.items()}
    assert not any(leaked.values()), (
        f"the raw literal is in the workbook the user downloads: {leaked!r}"
    )


@gap(
    "export:bi",
    "G5",
    "new",
    "audit-receipt-and-late-added-cards-skip-mask",
)
def test_bi_export_masks_the_audit_receipt_like_sql_used(harness: Any) -> None:
    """The BI stub's Cover table (no rows, no values) tokenises sql_used and ships the receipt raw."""
    _control_pack(harness)

    envelope = _receipt_envelope(NRIC_RECEIPT)
    envelope.update(rows=[], values=[])
    r = harness.client.post("/v1/chat/export.bi", json={"envelope": envelope})
    control(r.status_code == 200, f"export.bi: HTTP {r.status_code} {r.text[:200]}")
    body = r.json()
    control(body.get("source_table") == "cover", f"expected the Cover table: {body!r}")
    fields = {rec["field"]: rec["value"] for rec in body["table"]}
    control(
        "DMSMASK_nric" in fields.get("sql_used", "") and NRIC_RECEIPT not in fields["sql_used"],
        f"Cover sql_used was not masked, so the fixture is wrong: {fields.get('sql_used')!r}",
    )
    control("audit_receipt" in fields, f"Cover table has no audit_receipt: {sorted(fields)!r}")

    assert NRIC_RECEIPT not in r.text, (
        "the raw literal is in the BI stub (Power Query M, Superset dataset and Cover table): "
        f"{fields['audit_receipt']!r}"
    )


# --------------------------------------------------------------------------
# gap 2: the ledger append carries the raw SQL the envelope masks
# --------------------------------------------------------------------------

VQ_QUESTION = f"Payments for IC {NRIC_LEDGER}"
VQ_SQL = f"SELECT txn_id, quantity_kg FROM transactions WHERE txn_id <> '{NRIC_LEDGER}'"


def _control_ledger_append(
    h: Any, env: dict[str, Any], event_type: str, sql: str
) -> dict[str, Any]:
    """Controls shared by the ledger tests. Returns the payload DMS handed to Cortex."""
    require_envelope(200, env)
    control(
        env["abstained"] is False and env["rows"],
        f"the answer did not run: {env.get('text')!r}",
    )
    # The masker ran on this answer: sql_used is a token, the literal is not in it.
    control(
        "DMSMASK_nric" in str(env["sql_used"]) and NRIC_LEDGER not in str(env["sql_used"]),
        f"envelope sql_used was not masked, so the fixture is wrong: {env['sql_used']!r}",
    )
    # Cortex was handed the real SQL to execute (that is required), then one ledger entry.
    control(
        h.cortex.submitted_sql()[-1:] == [sql],
        f"Cortex was not handed the SQL: {h.cortex.submitted_sql()!r}",
    )
    payloads = h.cortex.ledger_payloads(event_type)
    control(len(payloads) == 1, f"expected one {event_type} append: {h.cortex.ledger_events()!r}")
    control(
        "sql" in payloads[0] and "run_id" in payloads[0],
        f"ledger payload lost its keys: {sorted(payloads[0])!r}",
    )
    return payloads[0]


@gap(
    "chat-ask:verified-query",
    "G5",
    "new",
    "ledger-append-carries-unmasked-sql",
)
def test_verified_query_ledger_payload_masks_sql_like_the_envelope(
    harness: Any, monkeypatch: Any
) -> None:
    """A steward VQ with a literal: sql_used is tokenised, the ledger entry holds the literal."""
    _control_pack(harness)
    _allow_studio_gate(monkeypatch)
    _register_verified_query(harness, VQ_QUESTION, VQ_SQL)

    status, env = harness.ask(VQ_QUESTION, space_id=FINANCE, session_id="gm_vq")
    control(status == 200 and env.get("badge") == "L0_CERTIFIED", f"{status}: {env!r}")
    payload = _control_ledger_append(harness, env, "ask.verified_query", VQ_SQL)

    assert NRIC_LEDGER not in json.dumps(payload), (
        f"the ledger entry holds the literal the envelope masks: {payload['sql']!r}"
    )


@gap(
    "mcp:ask",
    "G5",
    "new",
    "ledger-append-carries-unmasked-sql",
)
def test_mcp_ask_ledger_payload_masks_sql_like_the_envelope(
    harness_factory: Any, monkeypatch: Any
) -> None:
    """The same append through the MCP ask tool (the same handler as POST /v1/chat/ask)."""
    h = harness_factory(env={"DMS_MCP": "1"})
    _control_pack(h)
    _allow_studio_gate(monkeypatch)
    _register_verified_query(h, VQ_QUESTION, VQ_SQL)

    r = h.client.post(
        "/v1/mcp/call",
        json={
            "name": "ask",
            "arguments": {"question": VQ_QUESTION, "space_id": FINANCE, "session_id": "gm_mcp"},
        },
    )
    control(r.status_code == 200, f"mcp ask: HTTP {r.status_code} {r.text[:300]}")
    body = r.json()
    control(body.get("ok") is True and isinstance(body.get("result"), dict), repr(body)[:300])
    env = body["result"]
    control(env.get("badge") == "L0_CERTIFIED", f"{env!r}")
    payload = _control_ledger_append(h, env, "ask.verified_query", VQ_SQL)

    assert NRIC_LEDGER not in json.dumps(payload), (
        f"the ledger entry holds the literal the envelope masks: {payload['sql']!r}"
    )


@gap(
    "chat-ask:generative-insights-sql",
    "G5",
    "new",
    "ledger-append-carries-unmasked-sql",
)
def test_generative_insights_sql_ledger_payload_masks_sql_like_the_envelope(
    harness: Any,
) -> None:
    """A model-written SELECT that lifted a literal from the question: same split, generative lane."""
    sql = f"SELECT sku, quantity_kg FROM inventory WHERE sku <> '{NRIC_LEDGER}'"
    _control_pack(harness)
    # The Insights planner is a Cortex call and the harness fake has none, so a stand-in
    # returns the typed ``query_sql`` payload the live planner returns (no key, no network).
    harness.cortex.compute_insights = lambda question, **_kw: {"query_sql": sql}

    status, env = harness.ask(
        f"Show stock for every sku other than {NRIC_LEDGER}",
        space_id=FINANCE,
        session_id="gm_gen",
    )
    control(
        status == 200 and env.get("badge") == "L2_VALIDATED" and env.get("route") == "generated",
        f"the generative lane did not answer: {status} {env!r}",
    )
    payload = _control_ledger_append(harness, env, "ask.generated_ontology", sql)

    assert NRIC_LEDGER not in json.dumps(payload), (
        f"the ledger entry holds the literal the envelope masks: {payload['sql']!r}"
    )


# --------------------------------------------------------------------------
# gap 3: engine error text is returned verbatim as HTTP detail.message
# --------------------------------------------------------------------------


@gap(
    "ui:exclusion-auto-confirm",
    "G5",
    "new",
    "cortex-error-detail-returned-unmasked",
)
def test_engine_error_detail_is_masked_like_an_answer(harness_factory: Any) -> None:
    """A Cortex failure that quotes a cell value reaches the screen as 'ask 502: ...<address>'."""
    from cortex_client.generated.errors import UnexpectedStatus

    question = "Top 5 selling SKUs by revenue"
    # What the engine says when it fails on a cell: the HTTP error body, as the real client
    # raises it (UnexpectedStatus carries the response body), classified by the real executor.
    cast_failure = UnexpectedStatus(
        500, f'{{"detail":"cast failed for value \'{EMAIL}\' in column note"}}'.encode()
    )
    h = harness_factory(
        ask_response=AskResponse(
            answer=f"No row matched {EMAIL}.",
            badge="abstain",
            abstained=True,
            audit_id="aud_gm_err_ctl",
            route="abstain",
        )
    )
    _control_pack(h)

    # Control A: an ANSWER that quotes the same address is masked. The masker knows this literal.
    status, env = h.ask(question, space_id=FINANCE, session_id="gm_err_a")
    require_envelope(status, env)
    control(
        EMAIL not in json.dumps(env) and "DMSMASK_email" in env["text"],
        f"the answer text was not masked, so the fixture is wrong: {env['text']!r}",
    )

    # Control B: an error body with nothing sensitive keeps its code and message.
    h.cortex.ask_raises = UnexpectedStatus(403, b"pool_mismatch: pool mismatch")
    r = h.client.post(
        "/v1/chat/ask", json={"question": question, "space_id": FINANCE, "session_id": "gm_err_b"}
    )
    detail = r.json().get("detail") if r.status_code != 200 else None
    control(
        r.status_code == 403
        and isinstance(detail, dict)
        and detail.get("code") == "pool_mismatch"
        and "pool_mismatch" in str(detail.get("message")),
        f"the clean error path changed: HTTP {r.status_code} {r.text[:300]}",
    )

    # The gap request: the same route, the engine's error text quotes a cell value.
    h.cortex.ask_raises = cast_failure
    r = h.client.post(
        "/v1/chat/ask", json={"question": question, "space_id": FINANCE, "session_id": "gm_err_c"}
    )
    detail = r.json().get("detail") if r.status_code != 200 else None
    control(
        r.status_code == 502
        and isinstance(detail, dict)
        and detail.get("code") == "submit_failed",
        f"the engine failure did not arrive as an AskServiceError 502: HTTP {r.status_code} "
        f"{r.text[:300]}",
    )

    assert EMAIL not in r.text, (
        f"the HTTP error the UI prints carries the raw cell value: {detail['message']!r}"
    )

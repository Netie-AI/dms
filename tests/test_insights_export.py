"""INSIGHTS-EXPORT-01 — Excel from a real ask envelope, no invented rows."""

from __future__ import annotations

import io
import zipfile

from dms_api.app import create_app
from dms_api.settings import get_settings
from dms_core.xlsx_export import (
    EnvelopeExportError,
    export_envelope_xlsx,
    refuse_envelope_export,
    xlsx_download_name,
)
from fastapi.testclient import TestClient
from openpyxl import load_workbook

FINANCE = "cccccccc-cccc-cccc-cccc-cccccccccccc"

_ENVELOPE = {
    "answer_id": "ans_export_two",
    "badge": "L2_VALIDATED",
    "text": "Two SKUs by qty.",
    "values": [
        {"id": "v_a", "value": 1, "unit": "ea", "label": "A"},
        {"id": "v_b", "value": 2, "unit": "ea", "label": "B"},
    ],
    "rows": [
        {"sku": "SKU-ALPHA", "qty": 1},
        {"sku": "SKU-BETA", "qty": 2},
    ],
    "as_of": "2026-09-13T00:00:00Z",
    "sql_used": "SELECT sku, qty FROM t",
    "audit_id": "aud_export_two",
    "assumptions": [],
    "contributing_sources": [],
}


def _sheet_rows(data: bytes, name: str) -> list[list[object]]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb[name]
    out: list[list[object]] = []
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        while cells and cells[-1] is None:
            cells.pop()
        out.append(cells)
    return out


def test_refuse_without_real_envelope() -> None:
    assert refuse_envelope_export(None) == "envelope_required"
    assert refuse_envelope_export({}) == "envelope_required"
    assert refuse_envelope_export({"rows": [{"sku": "A"}]}) == "envelope_required"
    assert refuse_envelope_export({"answer_id": "ans_x"}) == "envelope_required"
    assert (
        refuse_envelope_export({"answer_id": "ans_x", "badge": "GREEN"}) == "envelope_required"
    )
    assert refuse_envelope_export(_ENVELOPE) is None


def test_xlsx_name_from_answer_id_not_clock() -> None:
    assert xlsx_download_name("ans_live_9f3c") == "dms_answer_ans_live_9f3c.xlsx"
    assert xlsx_download_name("ans/live\\9f3c") == "dms_answer_ans_live_9f3c.xlsx"
    assert xlsx_download_name("") == "dms_answer_export.xlsx"


def test_export_copies_rows_and_does_not_invent() -> None:
    data, name = export_envelope_xlsx(_ENVELOPE)
    assert name == "dms_answer_ans_export_two.xlsx"
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert "xl/workbook.xml" in names
    assert "xl/calcChain.xml" not in names
    grid = _sheet_rows(data, "Rows")
    assert grid[0] == ["sku", "qty"]
    assert grid[1] == ["SKU-ALPHA", 1]
    assert grid[2] == ["SKU-BETA", 2]
    assert len(grid) == 3
    values = _sheet_rows(data, "Values")
    assert values[1][1] == 1
    assert values[2][1] == 2
    cover = _sheet_rows(data, "Cover")
    fields = {row[0]: row[1] for row in cover[1:] if row}
    assert fields["answer_id"] == "ans_export_two"
    assert fields["badge"] == "L2_VALIDATED"
    assert "total" not in {str(c).lower() for row in grid for c in row}


def test_abstain_export_has_cover_only() -> None:
    env = {
        "answer_id": "ans_ab",
        "badge": "ABSTAIN",
        "abstained": True,
        "text": "No proof.",
        "values": [],
        "rows": [],
    }
    data, _name = export_envelope_xlsx(env)
    wb = load_workbook(io.BytesIO(data), read_only=True)
    assert wb.sheetnames == ["Cover"]
    cover = _sheet_rows(data, "Cover")
    fields = {row[0]: row[1] for row in cover[1:] if row}
    assert fields["badge"] == "ABSTAIN"
    assert fields["text"] == "No proof."


def test_malformed_rows_refused() -> None:
    bad = dict(_ENVELOPE)
    bad["rows"] = ["not-a-dict"]
    try:
        export_envelope_xlsx(bad)
    except EnvelopeExportError as exc:
        assert exc.code == "envelope_required"
    else:
        raise AssertionError("expected EnvelopeExportError")


def test_http_ask_then_export_demo(monkeypatch) -> None:
    monkeypatch.setenv("DMS_ASK_MODE", "demo")
    monkeypatch.delenv("DATABASE_URL", raising=False)
    get_settings.cache_clear()
    client = TestClient(create_app())
    asked = client.post(
        "/v1/chat/ask",
        json={"question": "What was total revenue?", "space_id": FINANCE},
    )
    assert asked.status_code == 200, asked.text
    env = asked.json()
    assert env["badge"] == "L2_VALIDATED"
    assert env["rows"]
    exported = client.post("/v1/chat/export.xlsx", json={"envelope": env})
    assert exported.status_code == 200, exported.text
    assert exported.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "dms_answer_" in exported.headers.get("content-disposition", "")
    grid = _sheet_rows(exported.content, "Rows")
    assert grid[0] == list(env["rows"][0].keys())
    assert len(grid) - 1 == len(env["rows"])
    body_qty = [tuple(row.get(k) for k in grid[0]) for row in env["rows"]]
    file_qty = [tuple(row) for row in grid[1:]]
    assert file_qty == body_qty


def test_http_export_refuses_without_envelope() -> None:
    client = TestClient(create_app())
    missing = client.post("/v1/chat/export.xlsx", json={})
    assert missing.status_code == 422
    empty = client.post("/v1/chat/export.xlsx", json={"envelope": {}})
    assert empty.status_code == 400
    assert empty.json()["detail"]["code"] == "envelope_required"
    invented = client.post(
        "/v1/chat/export.xlsx",
        json={"envelope": {"rows": [{"sku": "FAKE", "qty": 99}]}},
    )
    assert invented.status_code == 400


def test_http_export_does_not_reask(monkeypatch) -> None:
    monkeypatch.setenv("DMS_ASK_MODE", "demo")
    get_settings.cache_clear()
    app = create_app()

    def boom(*_a, **_k):
        raise AssertionError("export must not re-ask")

    app.state.ask_service.live_ask = boom
    app.state.ask_service.demo_ask = boom
    client = TestClient(app)
    r = client.post("/v1/chat/export.xlsx", json={"envelope": _ENVELOPE})
    assert r.status_code == 200, r.text
    assert len(_sheet_rows(r.content, "Rows")) == 3
